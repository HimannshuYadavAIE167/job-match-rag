import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Target files in sequential project build order
FILES_TO_EXPORT = [
    "requirements.txt",
    ".gitignore",
    ".env.example",
    "README.md",
    "src/__init__.py",
    "src/scrape_jobs.py",
    "src/preprocess.py",
    "src/embed.py",
    "src/retrieve.py",
    "src/generate.py",
    "src/pipeline.py",
    "eval/generate_ground_truth.py",
    "eval/labeled_pairs.csv",
    "eval/evaluate.py",
    "app/streamlit_app.py",
    ".streamlit/config.toml"
]

OUTPUT_DOCX = "JobMatch_RAG_Complete_Codebase.docx"


def set_font(run, font_name="Calibri", font_size=Pt(10), color=RGBColor(30, 41, 59), bold=False):
    run.font.name = font_name
    run.font.size = font_size
    run.font.color.rgb = color
    run.bold = bold


def add_code_block(doc, code_text):
    """Formats code using a distinct monospaced font inside a tinted box."""
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.5)

    cell = table.cell(0, 0)
    # Light gray background padding
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F1F5F9"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

    # Set light border
    borders_elm = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="2563EB"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders_elm)

    cell_para = cell.paragraphs[0]
    cell_para.paragraph_format.space_before = Pt(4)
    cell_para.paragraph_format.space_after = Pt(4)
    cell_para.paragraph_format.line_spacing = 1.15

    lines = code_text.strip().split("\n")
    for idx, line in enumerate(lines):
        run = cell_para.add_run(line if line else " ")
        set_font(run, font_name="Consolas", font_size=Pt(8.5), color=RGBColor(15, 23, 42))
        if idx < len(lines) - 1:
            cell_para.add_run("\n")


def generate_codebase_docx():
    doc = Document()

    # Document Margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Cover / Header
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("JobMatch RAG: Production Source Code & Architecture")
    set_font(title_run, font_name="Calibri", font_size=Pt(20), color=RGBColor(37, 99, 235), bold=True)

    desc_p = doc.add_paragraph()
    desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_run = desc_p.add_run("Full implementation across scraping, preprocessing, ChromaDB embeddings, Gemini generation, evaluation, and Streamlit.")
    set_font(desc_run, font_name="Calibri", font_size=Pt(10), color=RGBColor(100, 116, 139))

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    total_files = 0
    for relative_path in FILES_TO_EXPORT:
        file_path = Path(relative_path)
        
        # Add heading
        h = doc.add_heading(level=2)
        h_run = h.add_run(f"📁 {relative_path}")
        set_font(h_run, font_name="Calibri", font_size=Pt(13), color=RGBColor(30, 41, 59), bold=True)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(4)

        if not file_path.exists():
            warn_p = doc.add_paragraph()
            w_run = warn_p.add_run(f"[File not found in local repository directory: {relative_path}]")
            set_font(w_run, font_name="Calibri", font_size=Pt(9), color=RGBColor(220, 38, 38))
            continue

        try:
            content = file_path.read_text(encoding="utf-8")
            if not content.strip():
                content = "# (Empty file)"
            add_code_block(doc, content)
            total_files += 1
        except Exception as e:
            err_p = doc.add_paragraph()
            e_run = err_p.add_run(f"[Error reading file: {e}]")
            set_font(e_run, font_name="Calibri", font_size=Pt(9), color=RGBColor(220, 38, 38))

    doc.save(OUTPUT_DOCX)
    print(f"\n[OK] Successfully compiled {total_files} files into: {OUTPUT_DOCX}")


if __name__ == "__main__":
    generate_codebase_docx()