"""
Extracts plain text from an uploaded resume file so it can be fed into the
retrieval + generation pipeline the same way pasted text is.

Supports .txt, .pdf, and .docx. Kept dependency-light and defensive: a
malformed or password-protected file should produce a clear error message
in the UI, not a stack trace.
"""
from __future__ import annotations

import io


class ResumeParseError(Exception):
    """Raised when a resume file can't be read or contains no extractable text."""


def extract_text_from_txt(file_bytes: bytes) -> str:
    try:
        return file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return file_bytes.decode("latin-1", errors="ignore")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(file_bytes))
    except Exception as e:
        raise ResumeParseError(f"Could not open PDF: {e}") from e

    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as e:
            raise ResumeParseError("This PDF is password-protected and can't be read.") from e

    pages_text = []
    for page in reader.pages:
        try:
            pages_text.append(page.extract_text() or "")
        except Exception:
            continue

    return "\n".join(pages_text)


def extract_text_from_docx(file_bytes: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ResumeParseError(f"Could not open DOCX: {e}") from e

    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)

    return "\n".join(parts)


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """
    Dispatches to the right extractor based on file extension and validates
    that something usable came out the other end.
    """
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    if ext == "txt":
        text = extract_text_from_txt(file_bytes)
    elif ext == "pdf":
        text = extract_text_from_pdf(file_bytes)
    elif ext == "docx":
        text = extract_text_from_docx(file_bytes)
    else:
        raise ResumeParseError(f"Unsupported file type: .{ext}. Please upload .txt, .pdf, or .docx.")

    text = text.strip()
    if not text:
        raise ResumeParseError(
            "No extractable text found in this file. If it's a scanned/image-based PDF, "
            "try pasting the resume text directly instead."
        )
    return text
